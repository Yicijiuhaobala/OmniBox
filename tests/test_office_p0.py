import tempfile
import unittest
from pathlib import Path

from backend.office import execute_office_local_action, preview_office_action


class OfficeP0Tests(unittest.TestCase):
    def _workbook(self, path: Path, changed=False):
        from openpyxl import Workbook

        workbook = Workbook()
        first = workbook.active
        first.title = "华东"
        first.append(["姓名", "金额"])
        first.append(["张三", 10 if not changed else 11])
        second = workbook.create_sheet("华南")
        second.append(["姓名", "金额"])
        second.append(["李四", 20])
        workbook.save(path)
        workbook.close()

    def test_excel_preview_split_merge_and_compare(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            source = root / "sales.xlsx"
            other = root / "sales-new.xlsx"
            self._workbook(source)
            self._workbook(other, changed=True)

            preview = preview_office_action(source, "excel_split", {})
            self.assertTrue(preview["preview"])
            self.assertIn("拆分为 2", preview["operations"][0])

            split = execute_office_local_action(source, str(root / "split"), "excel_split", {})
            self.assertEqual(len(split["outputs"]), 2)
            self.assertTrue(all(Path(item).exists() for item in split["outputs"]))

            merged = execute_office_local_action(source, str(root / "merged"), "excel_merge_sheets", {})
            self.assertTrue(Path(merged["output"]).exists())
            self.assertIn("2 个工作表", merged["result"])

            compared = execute_office_local_action(source, str(root / "compare"), "excel_compare", {"other_file": str(other)})
            self.assertTrue(Path(compared["output"]).exists())
            self.assertIn("1 项差异", compared["result"])

    def test_word_replace_is_previewed_before_copy_is_written(self):
        from docx import Document

        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            source = root / "notice.docx"
            document = Document()
            document.add_paragraph("旧公司名称与旧公司名称")
            document.save(source)
            preview = preview_office_action(source, "word_replace", {"old": "旧公司名称", "new": "新公司名称"})
            self.assertIn("预计替换 2 处", preview["operations"][0])
            result = execute_office_local_action(source, str(root / "output"), "word_replace", {"old": "旧公司名称", "new": "新公司名称"})
            self.assertTrue(Path(result["output"]).exists())
            self.assertIn("旧公司名称", "\n".join(paragraph.text for paragraph in Document(source).paragraphs))
            self.assertNotIn("旧公司名称", "\n".join(paragraph.text for paragraph in Document(result["output"]).paragraphs))


if __name__ == "__main__":
    unittest.main()
