import tempfile
import unittest
from pathlib import Path

from backend.presentation import (
    PresentationExtractRequest,
    deck_from_markdown,
    extract_presentation,
    inspect_deck,
    markdown_to_docx,
)


class PresentationTests(unittest.TestCase):
    def test_markdown_becomes_editable_slide_outline(self):
        deck = deck_from_markdown(
            "# 项目周报\n\n## 本周进展\n- 完成开发\n- 通过测试\n\n## 下周计划\n发布版本",
            max_slides=10,
        )
        self.assertEqual(deck["title"], "项目周报")
        self.assertEqual(len(deck["slides"]), 3)
        self.assertEqual(deck["slides"][1]["title"], "本周进展")
        self.assertEqual(deck["slides"][1]["body"], ["完成开发", "通过测试"])

    def test_docx_and_pptx_can_be_extracted(self):
        from docx import Document
        from pptx import Presentation

        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            docx_path = root / "report.docx"
            document = Document()
            document.add_heading("季度复盘", level=1)
            document.add_heading("核心结果", level=2)
            document.add_paragraph("收入稳步增长")
            document.save(docx_path)
            docx_deck = extract_presentation(PresentationExtractRequest(source=str(docx_path)))
            self.assertEqual(docx_deck["source_type"], "docx")
            self.assertIn("核心结果", [slide["title"] for slide in docx_deck["slides"]])

            pptx_path = root / "brief.pptx"
            presentation = Presentation()
            first = presentation.slides.add_slide(presentation.slide_layouts[1])
            first.shapes.title.text = "发布计划"
            first.placeholders[1].text = "灰度发布\n收集反馈"
            presentation.save(pptx_path)
            pptx_deck = extract_presentation(PresentationExtractRequest(source=str(pptx_path)))
            self.assertEqual(pptx_deck["source_type"], "pptx")
            self.assertEqual(pptx_deck["slides"][0]["title"], "发布计划")
            self.assertIn("灰度发布", pptx_deck["slides"][0]["body"][0])

    def test_markdown_exports_to_docx_without_overwriting(self):
        with tempfile.TemporaryDirectory() as directory:
            first = markdown_to_docx("# 标题\n\n- 项目一", "周报", directory)
            second = markdown_to_docx("# 标题\n\n- 项目二", "周报", directory)
            self.assertTrue(first.exists())
            self.assertTrue(second.exists())
            self.assertNotEqual(first, second)

    def test_long_markdown_is_split_on_semantic_punctuation(self):
        paragraph = "这是一个很长的段落。" * 30
        deck = deck_from_markdown(f"# 长文\n\n## 分析\n{paragraph}", max_slides=20)
        self.assertGreaterEqual(len(deck["slides"]), 2)
        self.assertGreater(len(deck["slides"][1]["body"]), 1)
        self.assertTrue(any("长段落" in item for item in deck["warnings"]))
        self.assertTrue(all(len(item) <= 150 for slide in deck["slides"] for item in slide["body"]))

    def test_excel_data_becomes_chart_and_table_slides(self):
        from openpyxl import Workbook

        with tempfile.TemporaryDirectory() as directory:
            source = Path(directory) / "monthly.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "经营数据"
            sheet.append(["月份", "收入", "成本"])
            sheet.append(["一月", 12, 8])
            sheet.append(["二月", 18, 10])
            sheet.append(["三月", 24, 13])
            workbook.save(source)
            workbook.close()
            deck = extract_presentation(PresentationExtractRequest(source=str(source)))
            self.assertEqual(deck["source_type"], "xlsx")
            self.assertEqual(deck["slides"][1]["chart"]["series"][0]["name"], "收入")
            self.assertEqual(deck["slides"][1]["table"]["headers"][:3], ["月份", "收入", "成本"])

    def test_quality_inspection_reports_density_font_color_and_image_resolution(self):
        report = inspect_deck({"slides": [{
            "title": "非常长的页面标题" * 5,
            "body": ["内容" * 80] * 7,
            "images": [{"pixel_width": 640, "pixel_height": 360}],
            "source_audit": {"min_font_pt": 10, "colors": [str(index) for index in range(8)]},
        }]})
        codes = {item["code"] for item in report["issues"]}
        self.assertTrue({"long_title", "too_many_points", "dense_content", "small_font", "many_colors", "low_resolution_image"}.issubset(codes))
        self.assertLess(report["score"], 100)


if __name__ == "__main__":
    unittest.main()
