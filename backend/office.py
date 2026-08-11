from __future__ import annotations

import json
import re
import shutil
import zipfile
from copy import copy
from datetime import datetime
from pathlib import Path


EXCEL_SUFFIXES = {".xlsx", ".xlsm"}
WORD_SUFFIXES = {".docx"}


def resolve_office_file(raw_path: str) -> Path:
    path = Path(raw_path).expanduser().resolve()
    if not path.is_file():
        raise ValueError(f"文件不存在：{path}")
    if path.suffix.lower() not in EXCEL_SUFFIXES | WORD_SUFFIXES:
        raise ValueError("仅支持 XLSX、XLSM 和 DOCX 文件")
    return path


def _unique_path(path: Path) -> Path:
    if not path.exists():
        return path
    for index in range(2, 10_000):
        candidate = path.with_name(f"{path.stem}-{index}{path.suffix}")
        if not candidate.exists():
            return candidate
    raise ValueError(f"无法为 {path.name} 生成不重复的文件名")


def _output_path(source: Path, raw_output_dir: str, label: str) -> Path:
    output_dir = Path(raw_output_dir).expanduser().resolve() if raw_output_dir.strip() else (
        source.parent / "OmniBox 输出" / f"{datetime.now():%Y%m%d-%H%M%S}-Office"
    )
    output_dir.mkdir(parents=True, exist_ok=True)
    return _unique_path(output_dir / f"{source.stem}-{label}{source.suffix}")


def office_context(source: Path) -> dict:
    suffix = source.suffix.lower()
    if suffix in EXCEL_SUFFIXES:
        from openpyxl import load_workbook

        workbook = load_workbook(source, read_only=True, data_only=False, keep_vba=suffix == ".xlsm")
        try:
            sheets = []
            for sheet in workbook.worksheets[:12]:
                max_row, max_column = sheet.max_row, sheet.max_column
                if max_row is None or max_column is None:
                    from openpyxl.utils.cell import range_boundaries

                    _, _, calculated_column, calculated_row = range_boundaries(sheet.calculate_dimension(force=True))
                    max_row, max_column = calculated_row, calculated_column
                rows = []
                for row in sheet.iter_rows(min_row=1, max_row=min(max_row, 8), max_col=min(max_column, 12), values_only=True):
                    rows.append([str(value)[:160] if value is not None else "" for value in row])
                sheets.append({
                    "name": sheet.title,
                    "rows": max_row,
                    "columns": max_column,
                    "sample": rows,
                })
            return {"type": "excel", "file": source.name, "sheets": sheets}
        finally:
            workbook.close()

    from docx import Document

    document = Document(source)
    paragraphs = [paragraph.text.strip()[:240] for paragraph in document.paragraphs if paragraph.text.strip()][:40]
    tables = [{"rows": len(table.rows), "columns": len(table.columns)} for table in document.tables[:20]]
    return {
        "type": "word",
        "file": source.name,
        "paragraph_count": len(document.paragraphs),
        "sample_paragraphs": paragraphs,
        "tables": tables,
    }


def parse_plan(raw_text: str) -> dict:
    text = raw_text.strip()
    if text.startswith("```"):
        text = re.sub(r"^```(?:json)?\s*|\s*```$", "", text, flags=re.IGNORECASE)
    start, end = text.find("{"), text.rfind("}")
    if start < 0 or end <= start:
        raise ValueError("模型没有返回有效的 JSON 操作计划")
    plan = json.loads(text[start:end + 1])
    operations = plan.get("operations")
    if not isinstance(operations, list) or not operations:
        raise ValueError("操作计划为空")
    if len(operations) > 20:
        raise ValueError("单次最多执行 20 个 Office 操作")
    return plan


def _selected_sheets(workbook, raw_name: str):
    name = str(raw_name or "*").strip()
    if name == "*":
        return list(workbook.worksheets)
    if name not in workbook.sheetnames:
        raise ValueError(f"工作表不存在：{name}")
    return [workbook[name]]


def _excel_trim_text(sheet) -> int:
    from openpyxl.cell.cell import MergedCell

    changed = 0
    for row in sheet.iter_rows():
        for cell in row:
            if isinstance(cell, MergedCell) or not isinstance(cell.value, str) or cell.value.startswith("="):
                continue
            cleaned = cell.value.strip()
            if cleaned != cell.value:
                cell.value = cleaned
                changed += 1
    return changed


def _excel_remove_blank_rows(sheet, start_row: int = 2) -> int:
    removed = 0
    for row_index in range(sheet.max_row, max(1, start_row) - 1, -1):
        if all(cell.value is None or (isinstance(cell.value, str) and not cell.value.strip()) for cell in sheet[row_index]):
            sheet.delete_rows(row_index)
            removed += 1
    return removed


def _excel_deduplicate(sheet, columns: list | None = None, start_row: int = 2) -> int:
    from openpyxl.utils.cell import column_index_from_string

    if columns:
        indexes = []
        for value in columns:
            if isinstance(value, int):
                indexes.append(value)
            elif re.fullmatch(r"[A-Za-z]{1,3}", str(value).strip()):
                indexes.append(column_index_from_string(str(value).strip().upper()))
            else:
                raise ValueError(f"无效的去重列：{value}")
    else:
        indexes = list(range(1, sheet.max_column + 1))
    if any(index < 1 or index > sheet.max_column for index in indexes):
        raise ValueError(f"去重列超出工作表 {sheet.title} 的范围")

    seen = set()
    duplicates = []
    for row_index in range(max(1, start_row), sheet.max_row + 1):
        key = tuple(sheet.cell(row_index, index).value for index in indexes)
        if key in seen:
            duplicates.append(row_index)
        else:
            seen.add(key)
    for row_index in reversed(duplicates):
        sheet.delete_rows(row_index)
    return len(duplicates)


def _excel_fill_formula(sheet, target_range: str, formula: str) -> int:
    from openpyxl.formula.translate import Translator
    from openpyxl.utils.cell import range_boundaries

    if not formula.startswith("=") or len(formula) > 2_000:
        raise ValueError("公式必须以 = 开头且不超过 2000 字符")
    if "[" in formula or re.search(r"\b(?:WEBSERVICE|RTD|DDE|HYPERLINK)\s*\(", formula, re.IGNORECASE):
        raise ValueError("不允许写入外部工作簿引用或联网公式")
    min_col, min_row, max_col, max_row = range_boundaries(target_range)
    if min_col != max_col or min_row < 1 or max_row - min_row > 100_000:
        raise ValueError("公式目标必须是单列且不超过 100001 行")
    origin = sheet.cell(min_row, min_col)
    for row_index in range(min_row, max_row + 1):
        cell = sheet.cell(row_index, min_col)
        cell.value = Translator(formula, origin=origin.coordinate).translate_formula(cell.coordinate)
    return max_row - min_row + 1


def _apply_excel_operations(workbook, operations: list[dict]) -> list[str]:
    reports = []
    allowed = {"trim_text", "remove_blank_rows", "deduplicate_rows", "fill_formula"}
    for operation in operations:
        operation_type = str(operation.get("type", ""))
        if operation_type not in allowed:
            raise ValueError(f"不支持的 Excel 操作：{operation_type}")
        sheets = _selected_sheets(workbook, operation.get("sheet", "*"))
        for sheet in sheets:
            if operation_type == "trim_text":
                count = _excel_trim_text(sheet)
                reports.append(f"{sheet.title}：清理 {count} 个文本单元格")
            elif operation_type == "remove_blank_rows":
                count = _excel_remove_blank_rows(sheet, int(operation.get("start_row", 2)))
                reports.append(f"{sheet.title}：删除 {count} 个空白行")
            elif operation_type == "deduplicate_rows":
                count = _excel_deduplicate(sheet, operation.get("columns"), int(operation.get("start_row", 2)))
                reports.append(f"{sheet.title}：删除 {count} 个重复行")
            else:
                count = _excel_fill_formula(sheet, str(operation.get("target_range", "")), str(operation.get("formula", "")))
                reports.append(f"{sheet.title}：写入 {count} 个公式")
    return reports


def _paragraph_has_visual(paragraph) -> bool:
    return bool(paragraph._p.xpath(".//w:drawing | .//w:object | .//w:pict | .//w:sectPr"))


def _trim_paragraph_runs(paragraph) -> int:
    nonempty = [run for run in paragraph.runs if run.text]
    if not nonempty:
        return 0
    before = "".join(run.text for run in nonempty)
    nonempty[0].text = nonempty[0].text.lstrip()
    nonempty[-1].text = nonempty[-1].text.rstrip()
    return int(before != "".join(run.text for run in nonempty))


def _remove_manual_page_breaks(document) -> int:
    removed = 0
    for paragraph in list(document.paragraphs):
        page_breaks = list(paragraph._p.xpath(".//w:br[@w:type='page']"))
        for line_break in page_breaks:
            line_break.getparent().remove(line_break)
            removed += 1
        if page_breaks and not paragraph.text.strip() and not _paragraph_has_visual(paragraph) and not paragraph._p.xpath(".//w:br"):
            paragraph._element.getparent().remove(paragraph._element)
    return removed


def _remove_empty_paragraphs(document) -> int:
    removed = 0
    for paragraph in list(document.paragraphs):
        if not paragraph.text.strip() and not _paragraph_has_visual(paragraph) and not paragraph._p.xpath(".//w:br"):
            paragraph._element.getparent().remove(paragraph._element)
            removed += 1
    return removed


def _remove_empty_table_rows(document) -> int:
    removed = 0
    for table in document.tables:
        for row in list(table.rows):
            if all(not cell.text.strip() for cell in row.cells):
                table._tbl.remove(row._tr)
                removed += 1
    return removed


def _replace_text(document, old: str, new: str) -> int:
    if not old:
        raise ValueError("替换文本不能为空")
    replaced = 0
    paragraphs = list(document.paragraphs)
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                paragraphs.extend(cell.paragraphs)
    for paragraph in paragraphs:
        for run in paragraph.runs:
            count = run.text.count(old)
            if count:
                run.text = run.text.replace(old, new)
                replaced += count
    return replaced


def _apply_word_operations(document, operations: list[dict]) -> list[str]:
    reports = []
    allowed = {"trim_whitespace", "remove_empty_paragraphs", "remove_manual_page_breaks", "remove_empty_table_rows", "replace_text"}
    for operation in operations:
        operation_type = str(operation.get("type", ""))
        if operation_type not in allowed:
            raise ValueError(f"不支持的 Word 操作：{operation_type}")
        if operation_type == "trim_whitespace":
            paragraphs = list(document.paragraphs)
            for table in document.tables:
                for row in table.rows:
                    for cell in row.cells:
                        paragraphs.extend(cell.paragraphs)
            count = sum(_trim_paragraph_runs(paragraph) for paragraph in paragraphs)
            reports.append(f"清理 {count} 个段落首尾空白")
        elif operation_type == "remove_empty_paragraphs":
            reports.append(f"删除 {_remove_empty_paragraphs(document)} 个空段落")
        elif operation_type == "remove_manual_page_breaks":
            reports.append(f"删除 {_remove_manual_page_breaks(document)} 个人工分页符")
        elif operation_type == "remove_empty_table_rows":
            reports.append(f"删除 {_remove_empty_table_rows(document)} 个空表格行")
        else:
            count = _replace_text(document, str(operation.get("old", "")), str(operation.get("new", "")))
            reports.append(f"完成 {count} 处文本替换")
    return reports


def clean_office(source: Path, output_dir: str, kind: str, options: dict) -> dict:
    if kind == "excel_clean":
        if source.suffix.lower() not in EXCEL_SUFFIXES:
            raise ValueError("Excel 清理仅支持 XLSX/XLSM 文件")
        operations = []
        if options.get("trim_text", True):
            operations.append({"type": "trim_text", "sheet": "*"})
        if options.get("remove_blank_rows", True):
            operations.append({"type": "remove_blank_rows", "sheet": "*", "start_row": 2})
        if options.get("deduplicate_rows", False):
            operations.append({"type": "deduplicate_rows", "sheet": "*", "start_row": 2})
        return execute_office_plan(source, output_dir, {"summary": "Excel 本地清理", "operations": operations})

    if source.suffix.lower() not in WORD_SUFFIXES:
        raise ValueError("Word 清理仅支持 DOCX 文件")
    operations = []
    if options.get("trim_whitespace", True):
        operations.append({"type": "trim_whitespace"})
    if options.get("remove_empty_paragraphs", True):
        operations.append({"type": "remove_empty_paragraphs"})
    if options.get("remove_manual_page_breaks", True):
        operations.append({"type": "remove_manual_page_breaks"})
    if options.get("remove_empty_table_rows", False):
        operations.append({"type": "remove_empty_table_rows"})
    return execute_office_plan(source, output_dir, {"summary": "Word 本地清理", "operations": operations})


def preview_office_action(source: Path, kind: str, options: dict) -> dict:
    suffix = source.suffix.lower()
    operations: list[str] = []
    warnings: list[str] = []
    if kind.startswith("excel_"):
        if suffix not in EXCEL_SUFFIXES:
            raise ValueError("该操作仅支持 XLSX/XLSM 工作簿")
        from openpyxl import load_workbook

        workbook = load_workbook(source, read_only=True, data_only=False, keep_vba=suffix == ".xlsm")
        try:
            if kind == "excel_clean":
                operations.append(f"扫描并清理 {len(workbook.sheetnames)} 个工作表中的文本空格和空白行")
                if options.get("deduplicate_rows"):
                    operations.append("按整行内容去重，保留每组首次出现的数据")
            elif kind == "excel_split":
                operations.append(f"拆分为 {len(workbook.sheetnames)} 个独立 XLSX 文件")
                operations.extend(f"生成：{sheet.title}.xlsx（{sheet.max_row or 0} 行）" for sheet in workbook.worksheets[:20])
            elif kind == "excel_merge_sheets":
                total_rows = sum(max(0, (sheet.max_row or 0) - 1) for sheet in workbook.worksheets)
                operations.append(f"把 {len(workbook.sheetnames)} 个工作表纵向合并为“合并结果”，预计 {total_rows} 行数据")
                operations.append("新增“来源工作表”列，并以第一个工作表的首行作为表头")
                warnings.append("不同工作表列结构不一致时，会按列位置合并并保留空值")
            elif kind == "excel_compare":
                other = resolve_office_file(str(options.get("other_file", "")))
                if other.suffix.lower() not in EXCEL_SUFFIXES:
                    raise ValueError("对比文件必须是 XLSX/XLSM")
                operations.append(f"逐单元格比较 {source.name} 与 {other.name}")
                operations.append("生成差异工作簿，列出工作表、单元格、原值和对比值")
            else:
                raise ValueError(f"不支持的 Excel 操作：{kind}")
        finally:
            workbook.close()
    else:
        if suffix not in WORD_SUFFIXES:
            raise ValueError("该操作仅支持 DOCX 文档")
        from docx import Document

        document = Document(source)
        if kind == "word_clean":
            operations.append(f"检查 {len(document.paragraphs)} 个正文段落和 {len(document.tables)} 个表格")
            operations.append("清理首尾空格、空段落和人工分页符，默认保留图片与分节设置")
        elif kind == "word_replace":
            old = str(options.get("old", ""))
            if not old:
                raise ValueError("请输入要查找的文字")
            paragraphs = list(document.paragraphs)
            for table in document.tables:
                for row in table.rows:
                    for cell in row.cells:
                        paragraphs.extend(cell.paragraphs)
            count = sum(run.text.count(old) for paragraph in paragraphs for run in paragraph.runs)
            operations.append(f"预计替换 {count} 处“{old[:60]}”")
            operations.append("结果另存为新 DOCX，不覆盖源文件")
        elif kind == "word_extract":
            with zipfile.ZipFile(source) as archive:
                images = [name for name in archive.namelist() if name.startswith("word/media/") and not name.endswith("/")]
            operations.append(f"提取 {len(document.tables)} 个表格到 XLSX")
            operations.append(f"提取 {len(images)} 张原始图片到独立文件夹")
        else:
            raise ValueError(f"不支持的 Word 操作：{kind}")
    return {"result": "执行前预览", "operations": operations, "warnings": warnings, "preview": True, "source_preserved": True}


def _copy_cell(source_cell, target_cell) -> None:
    target_cell.value = source_cell.value
    if source_cell.has_style:
        target_cell.font = copy(source_cell.font)
        target_cell.fill = copy(source_cell.fill)
        target_cell.border = copy(source_cell.border)
        target_cell.alignment = copy(source_cell.alignment)
        target_cell.number_format = source_cell.number_format
        target_cell.protection = copy(source_cell.protection)


def _excel_split(source: Path, output_dir: str) -> dict:
    from openpyxl import Workbook, load_workbook

    workbook = load_workbook(source, data_only=False, keep_vba=source.suffix.lower() == ".xlsm")
    directory = _output_path(source, output_dir, "拆分").with_suffix("")
    directory.mkdir(parents=True, exist_ok=True)
    outputs = []
    try:
        for source_sheet in workbook.worksheets:
            target_book = Workbook()
            target_sheet = target_book.active
            target_sheet.title = source_sheet.title
            for row in source_sheet.iter_rows():
                for cell in row:
                    _copy_cell(cell, target_sheet[cell.coordinate])
            for key, dimension in source_sheet.column_dimensions.items():
                target_sheet.column_dimensions[key].width = dimension.width
            safe_name = re.sub(r'[\\/:*?"<>|]', "_", source_sheet.title).strip() or "Sheet"
            output = _unique_path(directory / f"{safe_name}.xlsx")
            target_book.save(output)
            target_book.close()
            outputs.append(str(output))
    finally:
        workbook.close()
    return {"result": f"已拆分 {len(outputs)} 个工作表", "output": str(directory), "outputs": outputs, "operations": [f"生成 {Path(item).name}" for item in outputs], "source_preserved": True}


def _excel_merge_sheets(source: Path, output_dir: str) -> dict:
    from openpyxl import Workbook, load_workbook

    source_book = load_workbook(source, data_only=False, keep_vba=source.suffix.lower() == ".xlsm")
    target_book = Workbook()
    target = target_book.active
    target.title = "合并结果"
    written = 0
    sheet_count = len(source_book.sheetnames)
    try:
        first_sheet = source_book.worksheets[0]
        headers = [cell.value for cell in first_sheet[1]] if first_sheet.max_row else []
        target.append(["来源工作表", *headers])
        for sheet in source_book.worksheets:
            for row in sheet.iter_rows(min_row=2, values_only=True):
                target.append([sheet.title, *row])
                written += 1
        output = _output_path(source, output_dir, "合并工作表").with_suffix(".xlsx")
        target_book.save(output)
    finally:
        source_book.close()
        target_book.close()
    return {"result": f"已合并 {sheet_count} 个工作表", "output": str(output), "operations": [f"写入 {written} 行数据", "新增来源工作表列"], "source_preserved": True}


def _excel_compare(source: Path, output_dir: str, other_raw: str) -> dict:
    from openpyxl import Workbook, load_workbook

    other = resolve_office_file(other_raw)
    if other.suffix.lower() not in EXCEL_SUFFIXES:
        raise ValueError("对比文件必须是 XLSX/XLSM")
    left = load_workbook(source, read_only=True, data_only=False)
    right = load_workbook(other, read_only=True, data_only=False)
    result_book = Workbook()
    result = result_book.active
    result.title = "差异"
    result.append(["工作表", "单元格", source.name, other.name])
    differences = 0
    try:
        for sheet_name in sorted(set(left.sheetnames) | set(right.sheetnames)):
            left_sheet = left[sheet_name] if sheet_name in left.sheetnames else None
            right_sheet = right[sheet_name] if sheet_name in right.sheetnames else None
            max_row = min(max(left_sheet.max_row if left_sheet else 0, right_sheet.max_row if right_sheet else 0), 100_000)
            max_column = min(max(left_sheet.max_column if left_sheet else 0, right_sheet.max_column if right_sheet else 0), 2_000)
            if max_row * max_column > 2_000_000:
                raise ValueError(f"工作表 {sheet_name} 超过 200 万个待比较单元格，请先缩小数据范围")
            for row in range(1, max_row + 1):
                for column in range(1, max_column + 1):
                    left_value = left_sheet.cell(row, column).value if left_sheet else None
                    right_value = right_sheet.cell(row, column).value if right_sheet else None
                    if left_value != right_value:
                        coordinate = result.cell(row=1, column=column).column_letter + str(row)
                        result.append([sheet_name, coordinate, left_value, right_value])
                        differences += 1
                        if differences >= 200_000:
                            raise ValueError("差异超过 200000 项，请缩小对比范围")
        output = _output_path(source, output_dir, "差异对比").with_suffix(".xlsx")
        result_book.save(output)
    finally:
        left.close(); right.close(); result_book.close()
    return {"result": f"发现 {differences} 项差异", "output": str(output), "operations": [f"对比文件：{other.name}", f"记录 {differences} 项差异"], "source_preserved": True}


def _word_replace_action(source: Path, output_dir: str, options: dict) -> dict:
    old, new = str(options.get("old", "")), str(options.get("new", ""))
    if not old:
        raise ValueError("请输入要查找的文字")
    return execute_office_plan(source, output_dir, {"summary": "Word 批量替换", "operations": [{"type": "replace_text", "old": old, "new": new}]})


def _word_extract(source: Path, output_dir: str) -> dict:
    from docx import Document
    from openpyxl import Workbook

    base = _output_path(source, output_dir, "提取内容").with_suffix("")
    base.mkdir(parents=True, exist_ok=True)
    document = Document(source)
    outputs: list[str] = []
    if document.tables:
        workbook = Workbook()
        workbook.remove(workbook.active)
        for index, table in enumerate(document.tables, start=1):
            sheet = workbook.create_sheet(f"表格{index}")
            for row in table.rows:
                sheet.append([cell.text for cell in row.cells])
        table_output = base / "表格.xlsx"
        workbook.save(table_output)
        workbook.close()
        outputs.append(str(table_output))
    image_count = 0
    with zipfile.ZipFile(source) as archive:
        for member in archive.infolist():
            if not member.filename.startswith("word/media/") or member.is_dir():
                continue
            image_count += 1
            name = Path(member.filename).name
            target = _unique_path(base / name)
            with archive.open(member) as input_stream, target.open("wb") as output_stream:
                shutil.copyfileobj(input_stream, output_stream)
            outputs.append(str(target))
    return {"result": f"已提取 {len(document.tables)} 个表格和 {image_count} 张图片", "output": str(base), "outputs": outputs, "operations": [f"生成 {Path(item).name}" for item in outputs], "source_preserved": True}


def execute_office_local_action(source: Path, output_dir: str, kind: str, options: dict) -> dict:
    if kind in {"excel_clean", "word_clean"}:
        return clean_office(source, output_dir, kind, options)
    if kind == "excel_split":
        return _excel_split(source, output_dir)
    if kind == "excel_merge_sheets":
        return _excel_merge_sheets(source, output_dir)
    if kind == "excel_compare":
        return _excel_compare(source, output_dir, str(options.get("other_file", "")))
    if kind == "word_replace":
        return _word_replace_action(source, output_dir, options)
    if kind == "word_extract":
        return _word_extract(source, output_dir)
    raise ValueError(f"不支持的 Office 操作：{kind}")


def execute_office_plan(source: Path, output_dir: str, plan: dict) -> dict:
    operations = plan.get("operations") or []
    if not operations:
        raise ValueError("没有可执行的操作")
    suffix = source.suffix.lower()
    output = _output_path(source, output_dir, "已处理")
    if suffix in EXCEL_SUFFIXES:
        from openpyxl import load_workbook

        workbook = load_workbook(source, keep_vba=suffix == ".xlsm")
        try:
            reports = _apply_excel_operations(workbook, operations)
            workbook.save(output)
        finally:
            workbook.close()
    elif suffix in WORD_SUFFIXES:
        from docx import Document

        document = Document(source)
        reports = _apply_word_operations(document, operations)
        document.save(output)
    else:
        raise ValueError("仅支持 XLSX、XLSM 和 DOCX 文件")
    return {
        "result": plan.get("summary") or f"已执行 {len(operations)} 个操作",
        "output": str(output),
        "operations": reports,
        "source_preserved": True,
    }
