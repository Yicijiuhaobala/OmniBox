from __future__ import annotations

import base64
import re
import uuid
from datetime import datetime
from pathlib import Path
from typing import Any

from fastapi import APIRouter, HTTPException
from pydantic import BaseModel, Field, model_validator


router = APIRouter(prefix="/api/presentations", tags=["presentations"])

MAX_SOURCE_BYTES = 80 * 1024 * 1024
MAX_TEXT_BYTES = 2 * 1024 * 1024
MAX_EMBEDDED_IMAGE_BYTES = 12 * 1024 * 1024
SUPPORTED_SOURCES = {".md", ".markdown", ".txt", ".docx", ".pptx", ".xlsx", ".xlsm"}


class PresentationExtractRequest(BaseModel):
    source: str = Field(default="", max_length=4_096)
    markdown: str = Field(default="", max_length=500_000)
    title: str = Field(default="", max_length=160)
    max_slides: int = Field(default=30, ge=2, le=80)

    @model_validator(mode="after")
    def require_content(self):
        if not self.source.strip() and not self.markdown.strip():
            raise ValueError("请选择源文件或输入 Markdown 内容")
        return self


class MarkdownDocxRequest(BaseModel):
    markdown: str = Field(min_length=1, max_length=500_000)
    title: str = Field(default="Markdown 文档", max_length=160)
    output_dir: str = Field(default="", max_length=4_096)


class PresentationInspectRequest(BaseModel):
    deck: dict[str, Any]


def _safe_title(value: str, fallback: str = "未命名演示") -> str:
    cleaned = re.sub(r"\s+", " ", value).strip()
    return cleaned[:120] or fallback


def _safe_source(raw_path: str) -> Path:
    path = Path(raw_path).expanduser().resolve()
    if not path.is_file():
        raise ValueError(f"文件不存在：{path}")
    if path.suffix.lower() not in SUPPORTED_SOURCES:
        raise ValueError("仅支持 Markdown、TXT、DOCX、XLSX、XLSM 和 PPTX 文件")
    if path.stat().st_size > MAX_SOURCE_BYTES:
        raise ValueError("源文件不能超过 80 MB")
    return path


def _slide(
    title: str,
    body: list[str] | None = None,
    notes: str = "",
    images: list[dict[str, Any]] | None = None,
    chart: dict[str, Any] | None = None,
    table: dict[str, Any] | None = None,
    source_audit: dict[str, Any] | None = None,
) -> dict[str, Any]:
    return {
        "id": uuid.uuid4().hex[:12],
        "title": _safe_title(title, "继续"),
        "body": [re.sub(r"\s+", " ", item).strip()[:500] for item in (body or []) if item.strip()][:8],
        "notes": notes.strip()[:4_000],
        "images": images or [],
        "chart": chart,
        "table": table,
        "source_audit": source_audit or {},
    }


def _split_long_item(value: str) -> list[str]:
    compact = re.sub(r"\s+", " ", value).strip()
    if len(compact) <= 150:
        return [compact] if compact else []
    parts = [part.strip() for part in re.split(r"(?<=[。！？；!?;])", compact) if part.strip()]
    output: list[str] = []
    buffer = ""
    for part in parts or [compact]:
        if buffer and len(buffer) + len(part) > 150:
            output.append(buffer)
            buffer = part
        else:
            buffer += part
    if buffer:
        output.append(buffer)
    return output


def _chunk_sections(title: str, sections: list[tuple[str, list[str]]], max_slides: int) -> list[dict[str, Any]]:
    slides: list[dict[str, Any]] = [_slide(title, ["由 OmniBox 本地生成，可继续编辑大纲和视觉风格"])]
    for section_title, items in sections:
        compact = [part for item in items if item.strip() for part in _split_long_item(item)]
        if not compact:
            compact = ["本节暂无正文，可在大纲中补充"]
        for offset in range(0, len(compact), 6):
            suffix = "" if offset == 0 else "（续）"
            slides.append(_slide(f"{section_title}{suffix}", compact[offset:offset + 6]))
            if len(slides) >= max_slides:
                return slides
    return slides[:max_slides]


def deck_from_markdown(markdown: str, requested_title: str = "", max_slides: int = 30) -> dict[str, Any]:
    lines = markdown.replace("\r\n", "\n").replace("\r", "\n").split("\n")
    document_title = requested_title.strip()
    sections: list[tuple[str, list[str]]] = []
    current_title = "主要内容"
    current_items: list[str] = []
    in_code = False
    code_lines: list[str] = []

    def flush_code() -> None:
        nonlocal code_lines
        if code_lines:
            current_items.append("代码示例：" + " ".join(code_lines)[:420])
            code_lines = []

    def flush_section() -> None:
        nonlocal current_items
        flush_code()
        if current_items:
            sections.append((current_title, current_items))
            current_items = []

    for raw in lines:
        line = raw.strip()
        if line.startswith("```"):
            in_code = not in_code
            if not in_code:
                flush_code()
            continue
        if in_code:
            if line:
                code_lines.append(line)
            continue
        heading = re.match(r"^(#{1,6})\s+(.+)$", line)
        if heading:
            level, value = len(heading.group(1)), heading.group(2).strip()
            if level == 1 and not document_title:
                document_title = value
                continue
            flush_section()
            current_title = value
            continue
        if not line or line in {"---", "***", "___"}:
            continue
        line = re.sub(r"^[-*+]\s+", "", line)
        line = re.sub(r"^\d+[.)]\s+", "", line)
        line = re.sub(r"^>\s*", "", line)
        line = re.sub(r"!\[([^]]*)]\([^)]+\)", r"图片：\1", line)
        line = re.sub(r"\[([^]]+)\]\([^)]+\)", r"\1", line)
        line = re.sub(r"[*_`~]", "", line).strip()
        if line:
            current_items.append(line)
    flush_section()
    document_title = _safe_title(document_title, "Markdown 演示")
    if not sections:
        sections = [("主要内容", [document_title])]
    warnings = []
    if any(len(re.sub(r"\s+", " ", item)) > 150 for _, items in sections for item in items):
        warnings.append("长段落已按语义标点拆分，建议在大纲中复核分页")
    return {
        "title": document_title,
        "source_type": "markdown",
        "slides": _chunk_sections(document_title, sections, max_slides),
        "warnings": warnings,
    }


def _deck_from_docx(path: Path, requested_title: str, max_slides: int) -> dict[str, Any]:
    from docx import Document

    document = Document(path)
    title = requested_title.strip() or path.stem
    sections: list[tuple[str, list[str]]] = []
    current_title = "主要内容"
    current_items: list[str] = []
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue
        style_name = str(paragraph.style.name if paragraph.style else "").lower()
        is_heading = "heading" in style_name or "标题" in style_name
        if is_heading:
            if not sections and not current_items and ("title" in style_name or "标题 1" in style_name) and not requested_title:
                title = text
                continue
            if current_items:
                sections.append((current_title, current_items))
            current_title, current_items = text, []
        else:
            current_items.append(text)
    for index, table in enumerate(document.tables[:10], start=1):
        rows = []
        for row in table.rows[:8]:
            values = [cell.text.strip() for cell in row.cells]
            if any(values):
                rows.append(" | ".join(values)[:500])
        if rows:
            current_items.append(f"表格 {index}：" + "；".join(rows))
    if current_items:
        sections.append((current_title, current_items))
    title = _safe_title(title, path.stem)
    return {
        "title": title,
        "source_type": "docx",
        "slides": _chunk_sections(title, sections or [("主要内容", [title])], max_slides),
        "warnings": ["复杂分页、浮动对象和页眉页脚不会进入演示大纲"],
    }


def _picture_data(shape, image_budget: list[int]) -> dict[str, Any] | None:
    try:
        image = shape.image
        blob = image.blob
    except (AttributeError, ValueError):
        return None
    if image_budget[0] + len(blob) > MAX_EMBEDDED_IMAGE_BYTES:
        return None
    image_budget[0] += len(blob)
    extension = str(image.ext or "png").lower()
    mime = "image/jpeg" if extension in {"jpg", "jpeg"} else f"image/{extension}"
    pixel_width = pixel_height = 0
    try:
        pixel_width, pixel_height = image.size
    except (AttributeError, TypeError, ValueError):
        pass
    return {
        "name": f"image.{extension}",
        "data_url": f"data:{mime};base64,{base64.b64encode(blob).decode('ascii')}",
        "width": int(shape.width),
        "height": int(shape.height),
        "pixel_width": int(pixel_width),
        "pixel_height": int(pixel_height),
    }


def _source_slide_audit(source_slide) -> dict[str, Any]:
    font_sizes: list[float] = []
    colors: set[str] = set()
    for shape in source_slide.shapes:
        if not getattr(shape, "has_text_frame", False):
            continue
        for paragraph in shape.text_frame.paragraphs:
            for run in paragraph.runs:
                if run.font.size:
                    font_sizes.append(round(run.font.size.pt, 1))
                try:
                    if run.font.color.rgb:
                        colors.add(str(run.font.color.rgb))
                except (AttributeError, ValueError):
                    pass
    return {
        "min_font_pt": min(font_sizes) if font_sizes else None,
        "max_font_pt": max(font_sizes) if font_sizes else None,
        "font_samples": len(font_sizes),
        "colors": sorted(colors)[:20],
    }


def _deck_from_pptx(path: Path, requested_title: str, max_slides: int) -> dict[str, Any]:
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    presentation = Presentation(path)
    slides: list[dict[str, Any]] = []
    image_budget = [0]
    for number, source_slide in enumerate(presentation.slides, start=1):
        if len(slides) >= max_slides:
            break
        slide_title = ""
        body: list[str] = []
        images: list[dict[str, Any]] = []
        title_shape = source_slide.shapes.title
        if title_shape is not None and getattr(title_shape, "text", "").strip():
            slide_title = title_shape.text.strip()
        for shape in source_slide.shapes:
            if (title_shape is None or shape.shape_id != title_shape.shape_id) and getattr(shape, "has_text_frame", False):
                text = re.sub(r"\s+", " ", shape.text).strip()
                if text:
                    body.append(text)
            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE and len(images) < 3:
                picture = _picture_data(shape, image_budget)
                if picture:
                    images.append(picture)
        notes = ""
        try:
            if source_slide.has_notes_slide:
                notes = source_slide.notes_slide.notes_text_frame.text
        except (AttributeError, ValueError):
            pass
        slides.append(_slide(slide_title or f"第 {number} 页", body, notes, images, source_audit=_source_slide_audit(source_slide)))
    title = _safe_title(requested_title or (slides[0]["title"] if slides else path.stem), path.stem)
    warnings = ["PPTX 将提取文本、普通图片和演讲者备注，并按新模板重新排版"]
    if image_budget[0] >= MAX_EMBEDDED_IMAGE_BYTES:
        warnings.append("图片总量较大，超过 12 MB 的部分未载入预览")
    return {"title": title, "source_type": "pptx", "slides": slides or [_slide(title)], "warnings": warnings}


def _plain_cell(value: Any) -> str | float | int:
    if isinstance(value, bool):
        return str(value)
    if isinstance(value, (int, float)):
        return value
    return re.sub(r"\s+", " ", str(value or "")).strip()[:120]


def _deck_from_excel(path: Path, requested_title: str, max_slides: int) -> dict[str, Any]:
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True, keep_vba=path.suffix.lower() == ".xlsm")
    title = _safe_title(requested_title or path.stem, path.stem)
    slides = [_slide(title, ["由 Excel 数据自动生成图表和摘要，可继续编辑后导出 PPTX"])]
    warnings: list[str] = []
    try:
        worksheets = [sheet for sheet in workbook.worksheets if sheet.sheet_state == "visible"][:8]
        for sheet in worksheets:
            if len(slides) >= max_slides:
                break
            rows = []
            for row in sheet.iter_rows(min_row=1, max_row=min(sheet.max_row, 200), values_only=True):
                values = [_plain_cell(value) for value in row[:12]]
                if any(value != "" for value in values):
                    rows.append(values)
                if len(rows) >= 30:
                    break
            if not rows:
                continue
            width = max(len(row) for row in rows)
            rows = [row + [""] * (width - len(row)) for row in rows]
            headers = [str(value or f"列 {index + 1}") for index, value in enumerate(rows[0])]
            data_rows = rows[1:]
            categories = [str(row[0]) for row in data_rows[:12]]
            series = []
            for column in range(1, min(width, 7)):
                values = [row[column] for row in data_rows[:12]]
                numeric = [float(value) for value in values if isinstance(value, (int, float))]
                if len(numeric) >= max(2, len(values) // 2):
                    normalized = [float(value) if isinstance(value, (int, float)) else 0.0 for value in values]
                    series.append({"name": headers[column], "values": normalized})
            chart = {"type": "bar", "categories": categories, "series": series[:4]} if categories and series else None
            table = {"headers": headers[:8], "rows": [row[:8] for row in data_rows[:8]]}
            summary = [f"数据范围：{len(data_rows)} 行 × {width} 列"]
            if series:
                summary.append("已识别数值序列：" + "、".join(item["name"] for item in series[:4]))
            else:
                summary.append("未识别到连续数值列，已保留表格预览")
            slides.append(_slide(sheet.title, summary, chart=chart, table=table))
    finally:
        workbook.close()
    if len(slides) == 1:
        warnings.append("工作簿没有可生成演示的可见数据")
    else:
        warnings.append("图表基于工作表前 12 行数据自动生成，请核对表头、单位和统计口径")
    return {"title": title, "source_type": path.suffix.lower().lstrip("."), "slides": slides, "warnings": warnings}


def inspect_deck(deck: dict[str, Any]) -> dict[str, Any]:
    issues: list[dict[str, Any]] = []
    slides = deck.get("slides") if isinstance(deck.get("slides"), list) else []
    for index, slide in enumerate(slides[:80], start=1):
        title = str(slide.get("title") or "")
        body = [str(item) for item in slide.get("body", []) if str(item).strip()]
        characters = sum(len(item) for item in body)
        audit = slide.get("source_audit") if isinstance(slide.get("source_audit"), dict) else {}
        if len(title) > 32:
            issues.append({"severity": "warning", "page": index, "code": "long_title", "message": "标题超过 32 个字，建议缩短"})
        if len(body) > 6:
            issues.append({"severity": "warning", "page": index, "code": "too_many_points", "message": f"包含 {len(body)} 个要点，建议控制在 6 个以内"})
        if characters > 360:
            issues.append({"severity": "error", "page": index, "code": "dense_content", "message": f"正文约 {characters} 字，页面信息过密"})
        if any(len(item) > 100 for item in body):
            issues.append({"severity": "warning", "page": index, "code": "long_point", "message": "存在超过 100 字的单条要点"})
        minimum_font = audit.get("min_font_pt")
        if isinstance(minimum_font, (int, float)) and minimum_font < 16:
            issues.append({"severity": "error", "page": index, "code": "small_font", "message": f"源 PPT 最小字号为 {minimum_font:g} pt，投影阅读可能困难"})
        colors = audit.get("colors") if isinstance(audit.get("colors"), list) else []
        if len(colors) > 6:
            issues.append({"severity": "warning", "page": index, "code": "many_colors", "message": f"源页面使用 {len(colors)} 种文字颜色，建议统一配色"})
        for image in slide.get("images", [])[:3]:
            width, height = image.get("pixel_width", 0), image.get("pixel_height", 0)
            if isinstance(width, int) and isinstance(height, int) and width and height and (width < 800 or height < 450):
                issues.append({"severity": "warning", "page": index, "code": "low_resolution_image", "message": f"图片分辨率仅 {width}×{height}，全屏展示可能模糊"})
    errors = sum(item["severity"] == "error" for item in issues)
    warnings = sum(item["severity"] == "warning" for item in issues)
    return {
        "score": max(0, 100 - errors * 15 - warnings * 6),
        "summary": "未发现明显问题" if not issues else f"发现 {errors} 个严重问题、{warnings} 个改进建议",
        "issues": issues,
    }


def extract_presentation(payload: PresentationExtractRequest) -> dict[str, Any]:
    if payload.markdown.strip():
        return deck_from_markdown(payload.markdown, payload.title, payload.max_slides)
    path = _safe_source(payload.source)
    suffix = path.suffix.lower()
    if suffix in {".md", ".markdown", ".txt"}:
        if path.stat().st_size > MAX_TEXT_BYTES:
            raise ValueError("文本文件不能超过 2 MB")
        text = path.read_text(encoding="utf-8-sig")
        result = deck_from_markdown(text, payload.title or path.stem, payload.max_slides)
        result["source_type"] = suffix.lstrip(".")
        return result
    if suffix == ".docx":
        return _deck_from_docx(path, payload.title, payload.max_slides)
    if suffix == ".pptx":
        return _deck_from_pptx(path, payload.title, payload.max_slides)
    return _deck_from_excel(path, payload.title, payload.max_slides)


def markdown_to_docx(markdown: str, title: str, output_dir: str) -> Path:
    from docx import Document
    from docx.shared import Pt

    directory = Path(output_dir).expanduser().resolve() if output_dir.strip() else Path.home() / "Documents" / "OmniBox 输出"
    directory.mkdir(parents=True, exist_ok=True)
    safe_name = re.sub(r'[\\/:*?"<>|]', "_", _safe_title(title, "Markdown 文档"))
    candidate = directory / f"{safe_name}.docx"
    if candidate.exists():
        candidate = directory / f"{safe_name}-{datetime.now():%Y%m%d-%H%M%S}.docx"
    document = Document()
    for raw in markdown.replace("\r\n", "\n").split("\n"):
        line = raw.rstrip()
        heading = re.match(r"^(#{1,6})\s+(.+)$", line)
        if heading:
            document.add_heading(heading.group(2).strip(), level=min(len(heading.group(1)), 4))
        elif re.match(r"^[-*+]\s+", line):
            document.add_paragraph(re.sub(r"^[-*+]\s+", "", line), style="List Bullet")
        elif re.match(r"^\d+[.)]\s+", line):
            document.add_paragraph(re.sub(r"^\d+[.)]\s+", "", line), style="List Number")
        elif line.startswith(">"):
            paragraph = document.add_paragraph(line.lstrip("> "))
            paragraph.style = document.styles["Quote"]
        elif line:
            document.add_paragraph(re.sub(r"[*_`~]", "", line))
        else:
            document.add_paragraph("")
    document.styles["Normal"].font.size = Pt(10.5)
    document.save(candidate)
    return candidate


@router.post("/extract")
def presentation_extract(payload: PresentationExtractRequest) -> dict[str, Any]:
    try:
        return extract_presentation(payload)
    except (OSError, ValueError) as exc:
        raise HTTPException(400, f"演示内容提取失败：{exc}") from exc


@router.post("/markdown-docx")
def presentation_markdown_docx(payload: MarkdownDocxRequest) -> dict[str, Any]:
    try:
        output = markdown_to_docx(payload.markdown, payload.title, payload.output_dir)
        return {"result": "Word 文档已生成", "output": str(output), "source_preserved": True}
    except (OSError, ValueError) as exc:
        raise HTTPException(400, f"Word 导出失败：{exc}") from exc


@router.post("/inspect")
def presentation_inspect(payload: PresentationInspectRequest) -> dict[str, Any]:
    return inspect_deck(payload.deck)
