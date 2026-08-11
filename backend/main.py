from __future__ import annotations

import base64
import hashlib
import json
import mimetypes
import os
import re
import secrets
import shutil
import string
import subprocess
import time
import uuid
from csv import writer as csv_writer
from datetime import datetime
from pathlib import Path
from typing import Literal, Optional

import httpx
import uvicorn
from fastapi import FastAPI, HTTPException
from fastapi.middleware.cors import CORSMiddleware
from pydantic import BaseModel, Field

try:
    from .advanced import router as advanced_router
    from .config_store import get_secret, read_config, write_config
    from .fund_learning import router as fund_learning_router
    from .hybrid import router as hybrid_router
    from .learning import router as learning_router
    from .lan_transfer import router as lan_transfer_router
    from .presentation import router as presentation_router
    from .presentation_templates import router as presentation_templates_router
    from .office_plans import router as office_plans_router
    from .office import execute_office_local_action, execute_office_plan, office_context, parse_plan, preview_office_action, resolve_office_file
    from .system_tools import router as system_tools_router
except ImportError:
    from advanced import router as advanced_router
    from config_store import get_secret, read_config, write_config
    from fund_learning import router as fund_learning_router
    from hybrid import router as hybrid_router
    from learning import router as learning_router
    from lan_transfer import router as lan_transfer_router
    from presentation import router as presentation_router
    from presentation_templates import router as presentation_templates_router
    from office_plans import router as office_plans_router
    from office import execute_office_local_action, execute_office_plan, office_context, parse_plan, preview_office_action, resolve_office_file
    from system_tools import router as system_tools_router


APP_VERSION = "0.1.0"
API_VERSION = APP_VERSION

app = FastAPI(title="OmniBox Local API", version=APP_VERSION)
app.add_middleware(
    CORSMiddleware,
    allow_origins=["http://localhost:5173", "null"],
    allow_origin_regex=r"^https?://(?:localhost|127\.0\.0\.1):\d+$",
    allow_credentials=False,
    allow_methods=["*"],
    allow_headers=["*"],
)
app.include_router(advanced_router)
app.include_router(hybrid_router)
app.include_router(learning_router)
app.include_router(fund_learning_router)
app.include_router(system_tools_router)
app.include_router(lan_transfer_router)
app.include_router(presentation_router)
app.include_router(presentation_templates_router)
app.include_router(office_plans_router)


PROVIDERS = {
    "openai": {"base_url": "https://api.openai.com/v1", "model": "gpt-4.1-mini"},
    "deepseek": {"base_url": "https://api.deepseek.com/v1", "model": "deepseek-chat"},
    "siliconflow": {"base_url": "https://api.siliconflow.cn/v1", "model": "Qwen/Qwen3-8B"},
    "custom": {"base_url": "", "model": ""},
}

AI_PROMPTS = {
    "polish": "你是一位专业文字编辑。润色用户文字，保留原意，提升清晰度、逻辑和表达质感。只输出润色结果。",
    "summarize": "你是信息分析助手。提炼用户内容的核心结论、关键事实和行动项，使用清晰的层次。",
    "translate": "你是专业翻译。准确、自然地翻译用户内容，保留原格式与专有名词。",
    "extract": "你是结构化信息提取助手。从内容中提取人物、组织、日期、地点、数字、任务与风险，缺失项不要臆造。",
    "explain_code": "你是资深软件工程师。解释代码目的、执行流程、关键设计、潜在问题，并给出简洁改进建议。",
    "draft": "你是高质量写作助手。根据用户材料与要求输出可直接使用、结构清晰的成稿。",
    "code_naming": "你是资深软件工程师。根据业务语义生成简洁、明确、符合指定语言惯例的变量名或函数名；说明极简，只输出候选名称列表。",
    "markdown_continue": "你是克制的 Markdown 编辑。延续用户已有内容、语气和标题层级，只输出可直接追加的 Markdown 正文，不要解释。",
    "markdown_toc": "分析 Markdown 标题层级，生成可直接插入文档的目录。只输出 Markdown 目录，不要臆造不存在的标题。",
    "markdown_summary": "提炼 Markdown 文档的核心内容，输出简洁、忠实、可直接放在文首的摘要。只输出 Markdown 摘要。",
    "prompt_optimizer": (
        "你是一名提示词工程师。把用户给出的简单 Prompt 扩展为可直接使用、可维护的高质量提示词。"
        "使用 Markdown 输出以下固定结构：1. System Prompt；2. Context（使用清晰的可替换变量）；"
        "3. Few-Shot Examples（至少两个输入/输出示例）；4. User Prompt Template；5. 使用建议与检查清单。"
        "不要虚构用户未提供的业务事实，缺失信息请用 {{变量名}} 占位；保留用户原始目标，并补充输出格式、边界条件和失败处理。"
    ),
    "office_formula": (
        "你是资深 Excel 公式顾问。根据用户的白话需求和可选工作簿结构，输出可直接粘贴的 Excel 公式。"
        "优先选择 SUMIF/SUMIFS、XLOOKUP、INDEX+MATCH、FILTER、LET 等清晰且可审计的写法。"
        "必须说明公式应放在哪个单元格、向下填充方式、每个引用范围的含义，并给出不支持 XLOOKUP 时的兼容写法。"
        "不要声称已修改文件；若信息不足，明确列出需要用户确认的列、工作表或数据范围。"
    ),
    "office_plan": (
        "你是 Office 文件处理规划器，只能返回一个 JSON 对象，不要输出 Markdown 或解释。"
        "Excel 允许的操作：trim_text、remove_blank_rows、deduplicate_rows、fill_formula；"
        "Word 允许的操作：trim_whitespace、remove_empty_paragraphs、remove_manual_page_breaks、remove_empty_table_rows、replace_text。"
        "JSON 格式为 {\"summary\":\"简短说明\",\"operations\":[...]}。"
        "Excel 操作可包含 sheet（工作表名或 *）、start_row、columns（列字母数组）；fill_formula 必须包含 target_range 和 formula，"
        "公式从目标范围首格写起并自动相对填充。Word replace_text 必须包含 old 和 new。"
        "只能基于用户要求选择最少的必要操作，不得使用未列出的类型，不得要求执行脚本、宏、联网公式或外部工作簿引用。"
    ),
}


class BasicToolRequest(BaseModel):
    tool: Literal["text", "json", "base64", "hash", "generate", "timestamp"]
    action: str
    input: str = ""
    options: dict = Field(default_factory=dict)


class AIRequest(BaseModel):
    tool: Literal[
        "polish", "summarize", "translate", "extract", "explain_code", "draft",
        "code_naming", "markdown_continue", "markdown_toc", "markdown_summary", "prompt_optimizer",
    ]
    input: str = Field(min_length=1, max_length=100_000)
    instruction: str = Field(default="", max_length=2_000)


class SettingsUpdate(BaseModel):
    provider: Literal["openai", "deepseek", "siliconflow", "custom"]
    base_url: str
    model: str
    api_key: str = ""
    youdao_app_key: str = ""
    youdao_app_secret: str = ""
    baidu_app_id: str = ""
    baidu_app_key: str = ""


class FileSummaryRequest(BaseModel):
    files: list[str] = Field(min_length=1, max_length=30)
    instruction: str = Field(default="", max_length=2_000)
    output_dir: str = ""


class FileConvertRequest(BaseModel):
    files: list[str] = Field(min_length=1, max_length=100)
    target: Literal["txt", "csv", "xlsx", "docx", "pdf"]
    output_dir: str = ""


class FileRenameRequest(BaseModel):
    files: list[str] = Field(min_length=1, max_length=500)
    pattern: str = Field(min_length=1, max_length=160)
    start: int = Field(default=1, ge=0, le=999_999)
    keep_originals: bool = True
    output_dir: str = ""


class OCRAIRequest(BaseModel):
    files: list[str] = Field(min_length=1, max_length=8)
    output_format: Literal["text", "markdown_table", "json"] = "text"
    instruction: str = Field(default="", max_length=2_000)


class OfficeFormulaRequest(BaseModel):
    instruction: str = Field(min_length=1, max_length=4_000)
    file: str = ""


class OfficeProcessRequest(BaseModel):
    action: Literal["excel_clean", "excel_split", "excel_merge_sheets", "excel_compare", "word_clean", "word_replace", "word_extract", "ai_process"]
    file: str = Field(min_length=1, max_length=4_000)
    instruction: str = Field(default="", max_length=4_000)
    output_dir: str = ""
    options: dict = Field(default_factory=dict)
    preview_only: bool = False
    plan: Optional[dict] = None


def load_settings() -> dict:
    defaults = {"provider": "openai", **PROVIDERS["openai"]}
    return {**defaults, **read_config()}


def get_api_key() -> str:
    return get_secret("api_key")


def resolve_files(raw_files: list[str]) -> list[Path]:
    files = list(dict.fromkeys(Path(item).expanduser().resolve() for item in raw_files))
    missing = [str(item) for item in files if not item.is_file()]
    if missing:
        raise HTTPException(400, f"文件不存在：{missing[0]}")
    return files


def make_output_dir(files: list[Path], raw_output_dir: str, label: str) -> Path:
    if raw_output_dir.strip():
        output_dir = Path(raw_output_dir).expanduser().resolve()
    else:
        output_dir = files[0].parent / "OmniBox 输出" / f"{datetime.now():%Y%m%d-%H%M%S}-{label}"
    output_dir.mkdir(parents=True, exist_ok=True)
    return output_dir


def unique_path(path: Path) -> Path:
    if not path.exists():
        return path
    for index in range(2, 10_000):
        candidate = path.with_name(f"{path.stem}-{index}{path.suffix}")
        if not candidate.exists():
            return candidate
    raise ValueError(f"无法为 {path.name} 生成不重复的文件名")


def extract_document_text(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        from pypdf import PdfReader

        reader = PdfReader(str(path))
        return "\n\n".join((page.extract_text() or "").strip() for page in reader.pages).strip()
    if suffix == ".docx":
        from docx import Document

        document = Document(str(path))
        paragraphs = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
        for table in document.tables:
            paragraphs.extend("\t".join(cell.text for cell in row.cells) for row in table.rows)
        return "\n".join(paragraphs)
    if suffix == ".xlsx":
        from openpyxl import load_workbook

        workbook = load_workbook(path, read_only=True, data_only=True)
        blocks = []
        try:
            for sheet in workbook.worksheets:
                rows = []
                for row in sheet.iter_rows(values_only=True):
                    values = ["" if value is None else str(value) for value in row]
                    if any(values):
                        rows.append("\t".join(values))
                blocks.append(f"## 工作表：{sheet.title}\n" + "\n".join(rows))
        finally:
            workbook.close()
        return "\n\n".join(blocks)
    if suffix in {".txt", ".csv"}:
        return path.read_text(encoding="utf-8-sig")
    raise ValueError(f"不支持读取 {suffix or '无扩展名'} 文件")


async def request_ai(tool: str, text: str, extra_instruction: str = "") -> str:
    settings = load_settings()
    api_key = get_api_key()
    if not api_key:
        raise HTTPException(401, "请先在设置中配置 AI API Key")
    instruction = AI_PROMPTS[tool]
    if extra_instruction.strip():
        instruction += f"\n用户的额外要求：{extra_instruction.strip()}"
    body = {
        "model": settings["model"],
        "messages": [
            {"role": "system", "content": instruction},
            {"role": "user", "content": text},
        ],
        "temperature": 0.3,
    }
    endpoint = f"{settings['base_url'].rstrip('/')}/chat/completions"
    try:
        async with httpx.AsyncClient(timeout=120) as client:
            response = await client.post(endpoint, json=body, headers={"Authorization": f"Bearer {api_key}"})
            response.raise_for_status()
            data = response.json()
        return data["choices"][0]["message"]["content"]
    except httpx.HTTPStatusError as exc:
        detail = exc.response.text[:500]
        raise HTTPException(exc.response.status_code, f"AI 服务返回错误：{detail}") from exc
    except (httpx.HTTPError, KeyError, IndexError, json.JSONDecodeError) as exc:
        raise HTTPException(502, f"无法连接 AI 服务或响应格式异常：{exc}") from exc


async def request_ai_vision(payload: OCRAIRequest) -> str:
    settings = load_settings()
    api_key = get_api_key()
    if not api_key:
        raise HTTPException(401, "请先在设置中配置 AI API Key")

    files = resolve_files(payload.files)
    format_instruction = {
        "text": "按图片顺序准确提取文字，保留合理段落。",
        "markdown_table": "识别表格并输出 Markdown 表格；非表格内容按段落输出。",
        "json": "输出合法 JSON，包含 images 数组，每项含 file、text、tables；不要使用 Markdown 代码围栏。",
    }[payload.output_format]
    content = [{"type": "text", "text": f"{format_instruction}\n{payload.instruction}".strip()}]
    for path in files:
        if path.suffix.lower() not in {".png", ".jpg", ".jpeg", ".webp"}:
            raise HTTPException(400, f"AI OCR 暂不支持：{path.name}")
        mime = mimetypes.guess_type(path.name)[0] or "image/png"
        encoded = base64.b64encode(path.read_bytes()).decode("ascii")
        content.extend([
            {"type": "text", "text": f"文件：{path.name}"},
            {"type": "image_url", "image_url": {"url": f"data:{mime};base64,{encoded}"}},
        ])
    body = {
        "model": settings["model"],
        "messages": [
            {"role": "system", "content": "你是文档 OCR 助手。准确识别图片内容，纠正常见识别错误，不添加图片中不存在的信息。"},
            {"role": "user", "content": content},
        ],
        "temperature": 0.1,
    }
    endpoint = f"{settings['base_url'].rstrip('/')}/chat/completions"
    try:
        async with httpx.AsyncClient(timeout=180) as client:
            response = await client.post(endpoint, json=body, headers={"Authorization": f"Bearer {api_key}"})
            response.raise_for_status()
            data = response.json()
        return data["choices"][0]["message"]["content"]
    except httpx.HTTPStatusError as exc:
        raise HTTPException(exc.response.status_code, f"AI 服务返回错误：{exc.response.text[:500]}") from exc
    except (httpx.HTTPError, KeyError, IndexError, json.JSONDecodeError) as exc:
        raise HTTPException(502, f"无法连接 AI 服务或响应格式异常：{exc}") from exc


@app.get("/api/health")
def health() -> dict:
    return {"ok": True, "api_version": API_VERSION, "time": int(time.time())}


@app.get("/api/settings")
def read_settings() -> dict:
    settings = load_settings()
    api_key = get_api_key()
    youdao_key = get_secret("youdao_app_key")
    youdao_secret = get_secret("youdao_app_secret")
    baidu_app_id = get_secret("baidu_app_id")
    baidu_app_key = get_secret("baidu_app_key")
    return {
        **settings,
        "api_key": api_key,
        "youdao_app_key": youdao_key,
        "youdao_app_secret": youdao_secret,
        "baidu_app_id": baidu_app_id,
        "baidu_app_key": baidu_app_key,
        "has_api_key": bool(api_key),
        "has_youdao_credentials": bool(youdao_key and youdao_secret),
        "has_baidu_credentials": bool(baidu_app_id and baidu_app_key),
    }


@app.put("/api/settings")
def save_settings(payload: SettingsUpdate) -> dict:
    base_url = payload.base_url.strip().rstrip("/")
    model = payload.model.strip()
    if not base_url.startswith(("http://", "https://")):
        raise HTTPException(400, "API 地址必须以 http:// 或 https:// 开头")
    if not model:
        raise HTTPException(400, "模型名称不能为空")

    api_key = payload.api_key.strip()
    youdao_key = payload.youdao_app_key.strip()
    youdao_secret = payload.youdao_app_secret.strip()
    baidu_app_id = payload.baidu_app_id.strip()
    baidu_app_key = payload.baidu_app_key.strip()
    if bool(youdao_key) != bool(youdao_secret):
        raise HTTPException(400, "有道 App Key 和 App Secret 需要同时填写或同时清空")
    if bool(baidu_app_id) != bool(baidu_app_key):
        raise HTTPException(400, "百度 APP ID 和密钥需要同时填写或同时清空")

    write_config({
        **read_config(),
        "provider": payload.provider,
        "base_url": base_url,
        "model": model,
        "api_key": api_key,
        "youdao_app_key": youdao_key,
        "youdao_app_secret": youdao_secret,
        "baidu_app_id": baidu_app_id,
        "baidu_app_key": baidu_app_key,
    })
    return {
        "ok": True,
        "has_api_key": bool(api_key),
        "has_youdao_credentials": bool(youdao_key and youdao_secret),
        "has_baidu_credentials": bool(baidu_app_id and baidu_app_key),
    }


@app.post("/api/tools/basic")
def run_basic_tool(payload: BasicToolRequest) -> dict:
    text = payload.input
    action = payload.action

    try:
        if payload.tool == "text":
            if action == "clean":
                result = "\n".join(line.strip() for line in text.splitlines() if line.strip())
            elif action == "dedupe":
                result = "\n".join(dict.fromkeys(line for line in text.splitlines() if line.strip()))
            elif action == "upper":
                result = text.upper()
            elif action == "lower":
                result = text.lower()
            elif action == "stats":
                result = json.dumps({
                    "字符数": len(text),
                    "非空字符": len(re.sub(r"\s", "", text)),
                    "单词数": len(re.findall(r"\b[\w'-]+\b", text)),
                    "行数": len(text.splitlines()),
                }, ensure_ascii=False, indent=2)
            else:
                raise ValueError("不支持的文本操作")

        elif payload.tool == "json":
            parsed = json.loads(text)
            if action == "format":
                result = json.dumps(parsed, ensure_ascii=False, indent=2)
            elif action == "minify":
                result = json.dumps(parsed, ensure_ascii=False, separators=(",", ":"))
            else:
                raise ValueError("不支持的 JSON 操作")

        elif payload.tool == "base64":
            if action == "encode":
                result = base64.b64encode(text.encode("utf-8")).decode("ascii")
            elif action == "decode":
                result = base64.b64decode(text, validate=True).decode("utf-8")
            else:
                raise ValueError("不支持的 Base64 操作")

        elif payload.tool == "hash":
            algorithm = action if action in {"md5", "sha1", "sha256", "sha512"} else "sha256"
            result = hashlib.new(algorithm, text.encode("utf-8")).hexdigest()

        elif payload.tool == "generate":
            count = max(1, min(int(payload.options.get("count", 5)), 50))
            if action == "uuid":
                result = "\n".join(str(uuid.uuid4()) for _ in range(count))
            elif action == "password":
                length = max(8, min(int(payload.options.get("length", 20)), 128))
                alphabet = string.ascii_letters + string.digits + "!@#$%^&*_-+"
                result = "\n".join("".join(secrets.choice(alphabet) for _ in range(length)) for _ in range(count))
            else:
                raise ValueError("不支持的生成操作")

        elif payload.tool == "timestamp":
            if action == "now":
                result = str(int(time.time()))
            elif action == "milliseconds":
                result = str(int(time.time() * 1000))
            elif action == "to_date":
                stamp = float(text)
                if stamp > 10_000_000_000:
                    stamp /= 1000
                result = time.strftime("%Y-%m-%d %H:%M:%S", time.localtime(stamp))
            elif action == "to_timestamp":
                result = str(int(time.mktime(time.strptime(text.strip(), "%Y-%m-%d %H:%M:%S"))))
            else:
                raise ValueError("不支持的时间操作")
        else:
            raise ValueError("未知工具")
    except (ValueError, TypeError, json.JSONDecodeError, UnicodeDecodeError) as exc:
        raise HTTPException(400, f"处理失败：{exc}") from exc

    return {"result": result}


@app.post("/api/tools/ai")
async def run_ai_tool(payload: AIRequest) -> dict:
    return {"result": await request_ai(payload.tool, payload.input, payload.instruction)}


@app.post("/api/hybrid/ocr-ai")
async def ocr_ai(payload: OCRAIRequest) -> dict:
    return {"result": await request_ai_vision(payload)}


@app.post("/api/office/formula")
async def office_formula(payload: OfficeFormulaRequest) -> dict:
    context = None
    if payload.file.strip():
        try:
            source = resolve_office_file(payload.file)
            if source.suffix.lower() not in {".xlsx", ".xlsm"}:
                raise ValueError("公式生成只能附加 XLSX/XLSM 工作簿")
            context = office_context(source)
        except ValueError as exc:
            raise HTTPException(400, str(exc)) from exc
    request = {"需求": payload.instruction, "工作簿结构与样例": context or "未提供工作簿"}
    return {"result": await request_ai("office_formula", json.dumps(request, ensure_ascii=False))}


@app.post("/api/office/process")
async def office_process(payload: OfficeProcessRequest) -> dict:
    try:
        source = resolve_office_file(payload.file)
        if payload.action != "ai_process":
            if payload.preview_only:
                return preview_office_action(source, payload.action, payload.options)
            return execute_office_local_action(source, payload.output_dir, payload.action, payload.options)
        if not payload.instruction.strip():
            raise ValueError("请描述希望如何处理文件")
        if payload.plan:
            if not isinstance(payload.plan.get("operations"), list) or not payload.plan["operations"] or len(payload.plan["operations"]) > 20:
                raise ValueError("确认执行的操作计划无效")
            plan = payload.plan
        else:
            context = office_context(source)
            planning_input = json.dumps({"用户要求": payload.instruction, "文件结构与样例": context}, ensure_ascii=False)
            raw_plan = await request_ai("office_plan", planning_input)
            plan = parse_plan(raw_plan)
        if payload.preview_only:
            return {"result": plan.get("summary") or "AI 操作计划", "operations": plan["operations"], "plan": plan, "preview": True, "source_preserved": True}
        return {**execute_office_plan(source, payload.output_dir, plan), "plan": plan}
    except HTTPException:
        raise
    except (ValueError, json.JSONDecodeError) as exc:
        raise HTTPException(400, f"Office 处理失败：{exc}") from exc


@app.post("/api/files/summarize")
async def summarize_files(payload: FileSummaryRequest) -> dict:
    files = resolve_files(payload.files)
    output_dir = make_output_dir(files, payload.output_dir, "摘要")
    results = []
    failures = []
    for path in files:
        try:
            text = extract_document_text(path)
            if not text.strip():
                raise ValueError("没有提取到可总结的文字（扫描版 PDF 暂不支持 OCR）")
            if len(text) > 80_000:
                text = text[:80_000] + "\n\n[文档内容过长，已截取前 80000 字符]"
            summary = await request_ai(
                "summarize",
                f"文件名：{path.name}\n\n{text}",
                payload.instruction or "按文件输出核心结论、关键数据和行动项",
            )
            result_path = unique_path(output_dir / f"{path.stem}-摘要.md")
            result_path.write_text(f"# {path.name} 摘要\n\n{summary}\n", encoding="utf-8")
            results.append({"source": str(path), "output": str(result_path), "summary": summary})
        except HTTPException:
            raise
        except Exception as exc:
            failures.append({"source": str(path), "error": str(exc)})
    if not results:
        detail = failures[0]["error"] if failures else "没有可处理的文件"
        raise HTTPException(400, f"摘要生成失败：{detail}")
    combined = "\n\n".join(f"## {Path(item['source']).name}\n\n{item['summary']}" for item in results)
    return {"result": combined, "items": results, "failures": failures, "output_dir": str(output_dir)}


def convert_with_libreoffice(source: Path, target: str, output_dir: Path) -> Path:
    executable = shutil.which("soffice") or shutil.which("libreoffice")
    if not executable and os.name == "posix":
        mac_path = Path("/Applications/LibreOffice.app/Contents/MacOS/soffice")
        if mac_path.exists():
            executable = str(mac_path)
    if not executable:
        raise ValueError("转换为 PDF 需要先安装 LibreOffice")
    process = subprocess.run(
        [executable, "--headless", "--convert-to", target, "--outdir", str(output_dir), str(source)],
        capture_output=True,
        text=True,
        timeout=120,
        check=False,
    )
    generated = output_dir / f"{source.stem}.{target}"
    if process.returncode != 0 or not generated.exists():
        message = (process.stderr or process.stdout or "LibreOffice 转换失败").strip()
        raise ValueError(message[:500])
    return generated


def convert_file(source: Path, target: str, output_dir: Path) -> list[Path]:
    suffix = source.suffix.lower()
    if target == "txt" and suffix in {".pdf", ".docx", ".xlsx", ".csv"}:
        output = unique_path(output_dir / f"{source.stem}.txt")
        output.write_text(extract_document_text(source), encoding="utf-8")
        return [output]
    if target == "csv" and suffix == ".xlsx":
        from openpyxl import load_workbook

        workbook = load_workbook(source, read_only=True, data_only=True)
        outputs = []
        try:
            multiple = len(workbook.sheetnames) > 1
            for sheet in workbook.worksheets:
                safe_sheet = re.sub(r'[\\/:*?"<>|]', "_", sheet.title).strip() or "Sheet"
                stem = f"{source.stem}-{safe_sheet}" if multiple else source.stem
                output = unique_path(output_dir / f"{stem}.csv")
                with output.open("w", encoding="utf-8-sig", newline="") as stream:
                    csv = csv_writer(stream)
                    for row in sheet.iter_rows(values_only=True):
                        csv.writerow(["" if value is None else value for value in row])
                outputs.append(output)
        finally:
            workbook.close()
        return outputs
    if target == "xlsx" and suffix == ".csv":
        import csv
        from openpyxl import Workbook

        output = unique_path(output_dir / f"{source.stem}.xlsx")
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "数据"
        with source.open("r", encoding="utf-8-sig", newline="") as stream:
            for row in csv.reader(stream):
                sheet.append(row)
        workbook.save(output)
        return [output]
    if target == "docx" and suffix == ".txt":
        from docx import Document

        output = unique_path(output_dir / f"{source.stem}.docx")
        document = Document()
        for line in source.read_text(encoding="utf-8-sig").splitlines():
            document.add_paragraph(line)
        document.save(output)
        return [output]
    if target == "pdf" and suffix in {".docx", ".xlsx"}:
        expected = output_dir / f"{source.stem}.pdf"
        if expected.exists():
            expected = unique_path(expected)
            temporary_dir = output_dir / f".omnibox-{uuid.uuid4().hex}"
            temporary_dir.mkdir()
            try:
                generated = convert_with_libreoffice(source, target, temporary_dir)
                shutil.move(str(generated), expected)
            finally:
                shutil.rmtree(temporary_dir, ignore_errors=True)
            return [expected]
        return [convert_with_libreoffice(source, target, output_dir)]
    raise ValueError(f"不支持 {suffix or '无扩展名'} → {target.upper()}；请调整目标格式")


@app.post("/api/files/convert")
def convert_files(payload: FileConvertRequest) -> dict:
    files = resolve_files(payload.files)
    output_dir = make_output_dir(files, payload.output_dir, "转换")
    results = []
    failures = []
    for source in files:
        try:
            outputs = convert_file(source, payload.target, output_dir)
            results.append({"source": str(source), "outputs": [str(item) for item in outputs]})
        except Exception as exc:
            failures.append({"source": str(source), "error": str(exc)})
    if not results:
        detail = failures[0]["error"] if failures else "没有可处理的文件"
        raise HTTPException(400, f"转换失败：{detail}")
    message = f"已转换 {len(results)} 个文件，生成 {sum(len(item['outputs']) for item in results)} 个结果"
    if failures:
        message += f"；{len(failures)} 个文件未处理"
    return {"result": message, "items": results, "failures": failures, "output_dir": str(output_dir)}


def render_rename(pattern: str, source: Path, number: int) -> str:
    contains_extension = "{ext}" in pattern
    name = pattern.replace("{name}", source.stem).replace("{ext}", source.suffix.lstrip("."))
    name = re.sub(r"\{n(?::(\d+))?\}", lambda match: str(number).zfill(int(match.group(1) or 0)), name)
    name = re.sub(r'[\\/:*?"<>|]', "_", name).strip().strip(".")
    if not name:
        raise ValueError("重命名规则生成了空文件名")
    if contains_extension:
        return name
    return f"{name}{source.suffix}"


@app.post("/api/files/rename")
def rename_files(payload: FileRenameRequest) -> dict:
    files = resolve_files(payload.files)
    destinations = [source.with_name(render_rename(payload.pattern, source, payload.start + index)) for index, source in enumerate(files)]
    if len({str(item).lower() for item in destinations}) != len(destinations):
        raise HTTPException(400, "重命名规则产生了重复文件名，请加入 {n} 序号")
    results = []
    if payload.keep_originals:
        output_dir = make_output_dir(files, payload.output_dir, "重命名")
        for source, desired in zip(files, destinations):
            destination = unique_path(output_dir / desired.name)
            shutil.copy2(source, destination)
            results.append({"source": str(source), "output": str(destination)})
    else:
        output_dir = files[0].parent
        source_set = {str(item).lower() for item in files}
        conflicts = [item for item in destinations if item.exists() and str(item).lower() not in source_set]
        if conflicts:
            raise HTTPException(409, f"目标文件已存在：{conflicts[0].name}")
        staged = []
        try:
            for source in files:
                temporary = source.with_name(f".omnibox-rename-{uuid.uuid4().hex}{source.suffix}")
                source.rename(temporary)
                staged.append(temporary)
            for temporary, destination, source in zip(staged, destinations, files):
                temporary.rename(destination)
                results.append({"source": str(source), "output": str(destination)})
        except Exception:
            for temporary, source in zip(staged, files):
                if temporary.exists() and not source.exists():
                    temporary.rename(source)
            raise
    return {"result": f"已完成 {len(results)} 个文件的批量重命名", "items": results, "failures": [], "output_dir": str(output_dir)}


if __name__ == "__main__":
    port = int(os.getenv("OMNIBOX_API_PORT", "8000"))
    uvicorn.run(app, host="127.0.0.1", port=port, log_level="info")
